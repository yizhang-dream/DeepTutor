"""Tests for deeptutor.utils.document_images — embedded-image extraction.

Fixtures generate real DOCX/PPTX archives on the fly (same pattern as
``test_document_extractor``) with a synthetic PNG large enough to clear the
minimum-size budget.
"""

from __future__ import annotations

import io
import random

from docx import Document as DocxDocument
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PptxInches
import pytest

from deeptutor.utils.document_images import (
    build_marker,
    extract_docx_rich,
    extract_pdf_images,
    extract_pptx_rich,
    find_markers,
    image_index_from_name,
)


def _png_bytes(size: tuple[int, int] = (400, 300)) -> bytes:
    """A noisy PNG comfortably above the minimum-bytes budget."""
    rng = random.Random(11)
    image_bytes = bytes(rng.randrange(256) for _ in range(size[0] * size[1] * 3))
    from PIL import Image as PILImage

    image = PILImage.frombytes("RGB", size, image_bytes)
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture(scope="module")
def png() -> bytes:
    return _png_bytes()


def _docx_with_image(paragraphs: list[str], png_data: bytes, *, table: bool = False) -> bytes:
    doc = DocxDocument()
    doc.add_paragraph(paragraphs[0])
    doc.add_picture(io.BytesIO(png_data), width=Inches(2))
    for paragraph in paragraphs[1:]:
        doc.add_paragraph(paragraph)
    if table:
        table_rows = doc.add_table(rows=1, cols=2)
        table_rows.cell(0, 0).text = "cell-a"
        table_rows.cell(0, 1).text = "cell-b"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx_with_image(png_data: bytes) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(PptxInches(1), PptxInches(1), PptxInches(4), PptxInches(1))
    box.text_frame.text = "标题文本"
    slide.shapes.add_picture(io.BytesIO(png_data), PptxInches(1), PptxInches(2), PptxInches(3))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Markers
# ---------------------------------------------------------------------------


def test_marker_roundtrip() -> None:
    from deeptutor.utils.document_images import EmbeddedImage

    image = EmbeddedImage(name="image-03.png", mime_type="image/png", data=b"x")
    marker = build_marker(image)
    assert find_markers(f"前文\n{marker}\n后文") == [(3, "image-03.png")]


def test_image_index_from_name() -> None:
    assert image_index_from_name("image-01.png") == 1
    assert image_index_from_name("media/image-12.jpg") == 12
    assert image_index_from_name("nope.png") == 0


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def test_docx_rich_extracts_image_and_marker(png: bytes) -> None:
    data = _docx_with_image(["第一段", "第二段"], png)
    rich = extract_docx_rich(data)

    assert len(rich.collection.images) == 1
    image = rich.collection.images[0]
    assert image.mime_type == "image/png"
    assert image.data == png

    marker = build_marker(image)
    assert any(marker in paragraph for paragraph in rich.paragraphs)
    # Text order preserved: intro before the marker, outro after it.
    text = "\n".join(rich.paragraphs)
    assert text.index("第一段") < text.index(marker) < text.index("第二段")


def test_docx_rich_keeps_table_text(png: bytes) -> None:
    data = _docx_with_image(["开头"], png, table=True)
    rich = extract_docx_rich(data)
    joined = "\n".join(rich.paragraphs)
    assert "cell-a" in joined and "cell-b" in joined


def test_docx_rich_repeats_marker_for_reused_target(png: bytes) -> None:
    """The same image embedded twice is stored once but marked twice."""
    doc = DocxDocument()
    doc.add_picture(io.BytesIO(png), width=Inches(2))
    doc.add_paragraph("中间")
    doc.add_picture(io.BytesIO(png), width=Inches(2))
    buf = io.BytesIO()
    doc.save(buf)

    rich = extract_docx_rich(buf.getvalue())
    assert len(rich.collection.images) == 1
    markers = find_markers("\n".join(rich.paragraphs))
    assert len(markers) == 2


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------


def test_pptx_rich_extracts_image_and_slide_text(png: bytes) -> None:
    rich = extract_pptx_rich(_pptx_with_image(png))

    assert len(rich.slides) == 1
    assert "标题文本" in rich.slides[0]
    assert len(rich.collection.images) == 1
    assert build_marker(rich.collection.images[0]) in rich.slides[0]


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def test_pdf_images_extracts_and_maps_pages(tmp_path) -> None:
    pymupdf = pytest.importorskip("pymupdf")
    from PIL import Image as PILImage

    png_data = _png_bytes((300, 200))
    pdf_path = tmp_path / "figures.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_image(pymupdf.Rect(72, 72, 300, 220), stream=png_data)
    doc.save(str(pdf_path))
    doc.close()
    del PILImage

    result = extract_pdf_images(pdf_path.read_bytes())
    assert result.collection.images, "the inserted figure should be extracted"
    assert result.page_map
    page_number, indices = result.page_map[0]
    assert page_number == 1
    assert result.collection.images[indices[0]].data


def test_pdf_images_deduplicates_repeated_xref(tmp_path) -> None:
    pytest.importorskip("pymupdf")
    png_data = _png_bytes((300, 200))
    pdf_path = tmp_path / "banner.pdf"
    doc = __import__("pymupdf").open()
    for _ in range(3):
        page = doc.new_page(width=612, height=792)
        page.insert_image(__import__("pymupdf").Rect(72, 72, 300, 160), stream=png_data)
    doc.save(str(pdf_path))
    doc.close()

    result = extract_pdf_images(pdf_path.read_bytes())
    assert len(result.collection.images) == 1
    assert [page for page, _ in result.page_map] == [1]
